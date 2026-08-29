from io import BytesIO
from typing import Any

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib import colors

from pptx import Presentation
from pptx.util import Inches, Pt


def value_to_text(value: Any) -> str:
    """
    Convert nested AI output into readable text.
    """

    if value is None:
        return ""

    if isinstance(value, str):
        return value

    if isinstance(value, (int, float, bool)):
        return str(value)

    if isinstance(value, list):
        lines = []

        for index, item in enumerate(value):
            if isinstance(item, dict):
                lines.append(
                    dict_to_text(item)
                )
            else:
                lines.append(
                    f"• {value_to_text(item)}"
                )

        return "\n".join(lines)

    if isinstance(value, dict):
        return dict_to_text(value)

    return str(value)


def dict_to_text(data: dict) -> str:
    """
    Convert dictionary content to readable text.
    """

    parts = []

    for key, value in data.items():

        title = key.replace("_", " ").title()

        parts.append(title)

        converted = value_to_text(value)

        if converted:
            parts.append(converted)

        parts.append("")

    return "\n".join(parts)


def build_pdf(
    output_type: str,
    content: Any,
) -> BytesIO:

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.alignment = TA_CENTER

    heading_style = styles["Heading2"]
    body_style = styles["BodyText"]

    story = []

    # NEXUS heading
    story.append(
        Paragraph(
            "NEXUS",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "AI-Powered Cyber Intelligence Transformation",
            styles["Subtitle"],
        )
    )

    story.append(Spacer(1, 25))

    story.append(
        Paragraph(
            output_type.replace("_", " ").title(),
            heading_style,
        )
    )

    story.append(Spacer(1, 15))

    if isinstance(content, dict):

        for key, value in content.items():

            heading = key.replace(
                "_",
                " ",
            ).title()

            story.append(
                Paragraph(
                    heading,
                    heading_style,
                )
            )

            story.append(
                Spacer(1, 6)
            )

            if isinstance(value, list):

                for item in value:

                    text = value_to_text(item)

                    story.append(
                        Paragraph(
                            text.replace(
                                "\n",
                                "<br/>",
                            ),
                            body_style,
                        )
                    )

                    story.append(
                        Spacer(1, 6)
                    )

            else:

                text = value_to_text(value)

                story.append(
                    Paragraph(
                        text.replace(
                            "\n",
                            "<br/>",
                        ),
                        body_style,
                    )
                )

            story.append(
                Spacer(1, 12)
            )

    elif isinstance(content, list):

        for item in content:

            story.append(
                Paragraph(
                    value_to_text(item).replace(
                        "\n",
                        "<br/>",
                    ),
                    body_style,
                )
            )

            story.append(
                Spacer(1, 10)
            )

    else:

        story.append(
            Paragraph(
                value_to_text(content).replace(
                    "\n",
                    "<br/>",
                ),
                body_style,
            )
        )

    document.build(story)

    buffer.seek(0)

    return buffer


def build_presentation(
    content: list,
) -> BytesIO:

    presentation = Presentation()

    # Remove default first slide if necessary
    while len(presentation.slides) > 0:
        xml_slides = presentation.slides._sldIdLst
        xml_slides.remove(xml_slides[0])

    for index, slide_data in enumerate(content):

        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        title = slide.shapes.title

        title.text = slide_data.get(
            "slide_title",
            f"Slide {index + 1}",
        )

        title.text_frame.paragraphs[0].font.size = Pt(
            28
        )

        body = slide.placeholders[1]

        text_frame = body.text_frame

        text_frame.clear()

        bullet_points = slide_data.get(
            "bullet_points",
            [],
        )

        for point_index, point in enumerate(
            bullet_points
        ):

            if point_index == 0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()

            paragraph.text = str(point)
            paragraph.level = 0
            paragraph.font.size = Pt(20)

        speaker_notes = slide_data.get(
            "speaker_notes"
        )

        if speaker_notes:

            notes_slide = slide.notes_slide

            notes_text_frame = (
                notes_slide.notes_text_frame
            )

            notes_text_frame.text = (
                str(speaker_notes)
            )

    buffer = BytesIO()

    presentation.save(buffer)

    buffer.seek(0)

    return buffer